from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "logs" / "gTCN_model_report.docx"


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True)
        shade_cell(cell, "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    document.add_paragraph()
    return table


def add_bullets(document, items):
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_code(document, text):
    paragraph = document.add_paragraph()
    paragraph.style = "No Spacing"
    run = paragraph.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    return paragraph


def build_report():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 2"].font.name = "Aptos Display"

    title = document.add_heading("gTCN: Conditional Waveform Guitar Effect Model", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Technical report based on the implemented repository")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Version context: five-effect training run, Slurm job 14019")

    document.add_heading("1. Executive Summary", 1)
    document.add_paragraph(
        "gTCN is a causal, waveform-to-waveform neural renderer for guitar effects. "
        "It accepts a dry guitar waveform and a reference pair consisting of aligned "
        "dry and wet audio. The reference pair is converted into a 128-dimensional "
        "effect embedding z, while the dry input is converted into time-aligned "
        "performance features. A FiLM-conditioned dilated temporal convolutional "
        "network then predicts the wet waveform directly, without a spectrogram "
        "inversion or vocoder stage."
    )
    document.add_paragraph(
        "The current implementation trains one shared model jointly across five "
        "effects: Alpaca Rogue, BritPlexi, FenderDeluxe, FenderDouble, and minosaur. "
        "Fixed robust embeddings can be computed after training so subsequent "
        "inference can bypass the effect encoder."
    )

    document.add_heading("2. System Overview", 1)
    add_code(document, "dry guitar --------------------> PerformanceEncoder --> TCNRenderer --> predicted wet")
    add_code(document, "dry_ref + wet_ref (aligned) --> EffectEncoder -----> z --FiLM--> TCNRenderer")
    document.add_paragraph(
        "The model has three principal components. The PerformanceEncoder and "
        "TCNRenderer form the reusable backbone. The EffectEncoder is the "
        "effect-specific conditioning component and is the planned replacement "
        "boundary for a future hypernetwork."
    )
    add_table(
        document,
        ["Component", "Input", "Output", "Role"],
        [
            ["PerformanceEncoder", "(B, 1, T) dry waveform", "(B, 64, T)", "Preserves performance timing and waveform length"],
            ["EffectEncoder", "(B, 2, T_ref) reference pair", "(B, 128)", "Produces global effect embedding z"],
            ["TCNRenderer", "(B, 64, T), z", "(B, 1, T)", "Generates the wet waveform"],
        ],
    )

    document.add_heading("3. Performance Encoder", 1)
    document.add_paragraph(
        "The PerformanceEncoder processes only the dry input. It contains two "
        "one-dimensional convolutions with 15-sample kernels, 64 channels, ReLU "
        "activations, and causal left padding. There is no stride or pooling. "
        "Consequently, its output has the same number of time samples as its input."
    )
    add_table(
        document,
        ["Input duration at 44.1 kHz", "Input samples T", "Encoder output"],
        [
            ["1 second", "44,100", "(B, 64, 44,100)"],
            ["30 seconds", "1,323,000", "(B, 64, 1,323,000)"],
            ["3 minutes", "7,938,000", "(B, 64, 7,938,000)"],
        ],
    )
    document.add_paragraph(
        "The channel dimension is fixed at 64 by residual_channels in "
        "configs/default.yaml. The time dimension scales directly with input "
        "duration. This is useful for causal rendering but means full-clip "
        "inference consumes memory proportional to clip length."
    )

    document.add_heading("4. Effect Encoder and Reference Crops", 1)
    document.add_paragraph(
        "The EffectEncoder concatenates dry_ref and wet_ref along the channel "
        "dimension, producing a two-channel input. Its convolutional stack uses "
        "15-sample kernels and three stride-4 layers, reducing the temporal "
        "resolution by approximately 64 times. Adaptive average pooling then "
        "collapses time, and a linear projection produces z."
    )
    add_bullets(
        document,
        [
            "Training reference duration: 2.0 seconds, or 88,200 samples at 44.1 kHz.",
            "The dry and wet reference crops are aligned and come from the same effect pair.",
            "The reference crop is sampled independently of the target training crop.",
            "The EffectEncoder output is (B, 128), independent of reference duration because of adaptive pooling.",
            "At explicit inference, complete reference files are currently accepted, but a representative 2-second aligned crop matches training more closely.",
        ],
    )

    document.add_heading("5. Causal Dilated TCN Renderer", 1)
    document.add_paragraph(
        "The renderer contains four repeated blocks, each with ten FiLM-conditioned "
        "residual layers. Within each block, dilation cycles through 1, 2, 4, ..., "
        "512. Each layer uses a kernel size of 3 and a gated activation: tanh for "
        "the filter branch multiplied by sigmoid for the gate branch. Residual and "
        "skip convolutions combine the layers into the final waveform prediction."
    )
    add_table(
        document,
        ["Setting", "Value", "Meaning"],
        [
            ["Residual channels", "64", "Feature width through the residual stack"],
            ["Skip channels", "64", "Width of accumulated skip features"],
            ["Blocks / layers", "4 / 10", "40 dilated residual layers total"],
            ["Kernel size", "3", "Three-sample temporal filter per layer"],
            ["Condition dimension", "128", "Size of z used by FiLM"],
            ["Receptive field", "8,212 samples", "Approximately 186.2 ms at 44.1 kHz"],
        ],
    )
    document.add_paragraph(
        "FiLM conditioning transforms each layer's intermediate activations using "
        "scale and bias values computed from z. The same dry performance can thus "
        "be rendered with different effect identities without changing the backbone."
    )

    document.add_heading("6. Training Data and Objective", 1)
    document.add_paragraph(
        "The dataset pairs dry files under GTtimbre/data/X with wet files under "
        "GTtimbre/data/y/<effect>. A dry file named item.wav normally maps to "
        "item-<effect>.wav. The loader also supports an unambiguous alternative "
        "suffix, which is how Alpaca Rogue files named item-reverb.wav are paired."
    )
    add_bullets(
        document,
        [
            "Training target crop: 1.0 second (44,100 samples).",
            "Reference crop: 2.0 seconds (88,200 samples).",
            "Validation split: 10 percent per effect, using a fixed seed.",
            "Training is joint across all discovered effect directories.",
            "The epoch length is fixed at 200 batches with batch size 8.",
        ],
    )
    document.add_paragraph(
        "The objective combines multi-resolution STFT loss and waveform L1 loss. "
        "STFT loss averages spectral convergence and log-magnitude errors at FFT "
        "sizes 2048, 1024, and 512. With both configured weights equal to 1.0, "
        "the total loss is spectral convergence plus log magnitude plus time-domain L1."
    )
    add_table(
        document,
        ["Training setting", "Configured value"],
        [
            ["Epochs", "300"],
            ["Batch size", "8"],
            ["Learning rate", "0.0003"],
            ["Optimizer", "Adam"],
            ["Gradient clipping", "1.0"],
            ["Training run", "Slurm job 14019, completed successfully"],
        ],
    )

    document.add_heading("7. Fixed Robust Effect Embeddings", 1)
    document.add_paragraph(
        "After training, compute_effect_embeddings.py samples 128 reference crops "
        "for each effect by default. Each sample selects one available dry/wet pair "
        "with replacement and then selects a random 2-second position within that "
        "pair. It does not split every complete clip into non-overlapping segments."
    )
    document.add_paragraph(
        "For each effect, the script encodes all sampled references, finds the "
        "coordinate-wise median center, measures each embedding's distance from that "
        "center, removes the furthest 10 percent, and averages the remaining 90 "
        "percent. With 128 samples, the saved robust mean uses 115 inliers."
    )
    add_table(
        document,
        ["Effect", "Fixed embedding file"],
        [
            ["Alpaca Rogue", "checkpoints/Alpaca_Rogue_embeddings.pt"],
            ["BritPlexi", "checkpoints/BritPlexi_embeddings.pt"],
            ["FenderDeluxe", "checkpoints/FenderDeluxe_embeddings.pt"],
            ["FenderDouble", "checkpoints/FenderDouble_embeddings.pt"],
            ["minosaur", "checkpoints/minosaur_embeddings.pt"],
        ],
    )
    add_code(document, "python src/compute_effect_embeddings.py --checkpoint checkpoints/best.pt --output-dir checkpoints")
    document.add_paragraph(
        "Each file stores the source checkpoint path, effect name, 128-dimensional "
        "tensor z, and sampling metadata. These files must be regenerated whenever "
        "the model checkpoint changes."
    )

    document.add_heading("8. Inference Behavior", 1)
    document.add_paragraph(
        "For deterministic rendering, GuitarRenderer loads a fixed embedding and "
        "calls render_with_z. The effect encoder is bypassed. Alternatively, the "
        "renderer can sample a 2-second reference from the training data when given "
        "an effect name, or encode explicitly supplied dry/wet reference files."
    )
    add_code(document, "python src/inference.py \\")
    add_code(document, "  --checkpoint checkpoints/best.pt \\")
    add_code(document, "  --embeddings checkpoints/BritPlexi_embeddings.pt \\")
    add_code(document, "  --input path/to/new_dry.wav \\")
    add_code(document, "  --output path/to/predicted.wav --timbre BritPlexi")
    document.add_paragraph(
        "The fixed z has shape (1, 128) regardless of whether the dry input is 30 "
        "seconds or 3 minutes. The performance features and output waveform retain "
        "the full input time length. Computation and activation memory grow roughly "
        "sixfold when duration grows from 30 seconds to 3 minutes."
    )
    document.add_paragraph(
        "The current inference implementation processes a complete input waveform in "
        "one forward pass. Long recordings may therefore exceed available GPU memory. "
        "Because the renderer is causal and its receptive field is about 186 ms, an "
        "overlap-aware chunked inference implementation is the appropriate next step "
        "for robust multi-minute rendering."
    )

    document.add_heading("9. Real-Time Streaming Development", 1)
    document.add_paragraph(
        "The current model is causal but not stateful: every call pads and processes "
        "the supplied waveform as a complete tensor. A playable application should "
        "replace this full-clip path with streaming blocks while preserving the "
        "causal history between blocks. The fixed effect embedding z can remain in "
        "memory for the duration of a performance."
    )
    add_bullets(
        document,
        [
            "Add forward_stream and reset_stream methods to the PerformanceEncoder and TCNRenderer.",
            "Cache the previous 14 dry samples and intermediate feature samples for the two kernel-15 performance convolutions.",
            "For each renderer layer with dilation d, cache 2d past feature samples because its kernel size is 3.",
            "Use small audio blocks, initially 256 to 1,024 samples, balancing CPU/GPU efficiency against latency.",
            "Keep audio input/output callbacks separate from model inference by using a ring buffer and worker thread.",
            "Use inference mode, evaluation mode, pinned buffers, and benchmarked device placement before connecting a live instrument.",
        ],
    )
    document.add_paragraph(
        "A simpler prototype can use overlap processing: prepend the previous 8,212 "
        "samples to each block, run the existing model, and discard the warm-up "
        "output. This is easier to validate but recomputes context on every block. "
        "State caching avoids that repeated work and is the preferred production design."
    )
    document.add_paragraph(
        "The current receptive field is approximately 186 ms. This is causal history "
        "rather than explicit lookahead, but it still affects state size and startup "
        "warm-up. If live playing latency is too high, reducing the number of blocks "
        "or layers per block can lower compute cost at the expense of modeling memory."
    )

    document.add_heading("10. Future Temporal and Long-Memory Effects", 1)
    document.add_paragraph(
        "The current training set and demonstrated workflow focus on distortion and "
        "amp-style effects. These are predominantly short-memory nonlinear effects: "
        "the output at a sample depends mainly on the current input and a relatively "
        "short recent history. The present causal TCN is well matched to that case."
    )
    document.add_paragraph(
        "Delay and reverb introduce a different requirement. Their output can depend "
        "on audio from hundreds of milliseconds, several seconds, or longer in the "
        "case of reverberant tails. A 186 ms receptive field cannot reproduce a "
        "long delay or a full reverb tail from local convolution context alone."
    )
    add_table(
        document,
        ["Effect family", "Main challenge", "Recommended development"],
        [
            ["Distortion / waveshaping", "Short-memory nonlinear transformation", "Current TCN and fixed-z conditioning"],
            ["Delay / echo", "Explicit long delay and repeated feedback", "Add a differentiable delay-line or recurrent state with delay-time and feedback conditioning"],
            ["Reverb", "Long, dense, possibly multi-second decay", "Use a hybrid neural renderer with explicit state, multiscale temporal processing, or a dedicated convolution/reverb tail module"],
        ],
    )
    add_bullets(
        document,
        [
            "Collect aligned dry/wet data containing complete audible tails, not only short target crops.",
            "Use longer reference crops or separate effect descriptors for decay time, delay time, feedback, and wet/dry mix.",
            "Extend the streaming state independently of the short TCN receptive field so memory length matches the effect.",
            "Train with losses that evaluate tails and long-term energy decay, in addition to local STFT and waveform losses.",
            "Consider separate short-memory and long-memory branches whose outputs are mixed by the effect embedding.",
        ],
    )
    document.add_paragraph(
        "For delay, an explicit circular buffer is usually more reliable and efficient "
        "than asking a convolutional network to learn a long exact delay. For reverb, "
        "a hybrid architecture can preserve a neural timbral transformation while an "
        "explicit state-space, feedback, or convolutional tail handles long decay. "
        "These changes should be introduced with new effect-specific evaluation rather "
        "than assuming the distortion-trained model will generalize automatically."
    )

    document.add_heading("11. Checkpoints and Outputs", 1)
    add_bullets(
        document,
        [
            "checkpoints/last.pt: latest completed training epoch.",
            "checkpoints/best.pt: best overall validation checkpoint.",
            "checkpoints/best_<effect>.pt: best checkpoint according to each effect's validation loss.",
            "checkpoints/<effect>_embeddings.pt: fixed robust z files generated after training.",
            "outputs/eval/: generated evaluation WAV files; excluded from Git tracking.",
            "outputs/tensorboard/: TensorBoard event files; excluded from Git tracking.",
        ],
    )

    document.add_heading("12. Design Strengths and Current Limitations", 1)
    document.add_paragraph("Strengths")
    add_bullets(
        document,
        [
            "Direct waveform synthesis avoids phase loss from Griffin-Lim reconstruction.",
            "Causal dilated convolutions provide a real-time-compatible local temporal structure.",
            "One shared backbone learns across multiple effects and supports effect conditioning.",
            "Robust fixed z reduces inference variability caused by a single reference crop.",
        ],
    )
    document.add_paragraph("Limitations and follow-up work")
    add_bullets(
        document,
        [
            "Full-clip inference currently holds the entire activation sequence in memory.",
            "The fixed embedding files are tied to the exact checkpoint that produced them.",
            "Explicit reference-file inference accepts variable lengths, although training uses 2-second crops.",
            "Evaluation currently focuses on waveform and spectral comparisons; perceptual listening tests remain useful.",
            "Chunked overlap-aware inference would make multi-minute input handling more reliable.",
        ],
    )

    document.add_heading("13. Repository Entry Points", 1)
    add_table(
        document,
        ["Path", "Purpose"],
        [
            ["src/model.py", "PerformanceEncoder, EffectEncoder, TCNRenderer, and GTCNModel"],
            ["src/dataset.py", "Dry/wet indexing, random crops, and multi-effect batches"],
            ["src/losses.py", "Multi-resolution STFT and time-domain loss"],
            ["src/train.py", "Training, validation, and checkpointing"],
            ["src/inference.py", "Rendering with references or fixed embeddings"],
            ["src/compute_effect_embeddings.py", "Robust fixed z generation"],
            ["configs/default.yaml", "Data, architecture, and training configuration"],
            ["notebooks/evaluate_fixed_z.ipynb", "Fixed-z audio and spectral evaluation"],
        ],
    )

    document.add_paragraph("Generated for the gTCN repository.").italic = True
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()